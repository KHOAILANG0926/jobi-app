#!/usr/bin/env python3
"""Extract the official Vietnamese occupation classification (QĐ 34/2020/QĐ-TTg).

Input : data/occupations/34-2020/source/*.doc  (official gazette files, see MANIFEST.json)
Output: data/occupations/34-2020/
          occupation_codes.json   Phụ lục I  — code, level, parent, name (5 levels)
          descriptions.json       Phụ lục II — description, main tasks, examples, exclusions
          needs_review.json       rows that could not be read unambiguously (never guessed)
          VALIDATION.json         counts + integrity checks

Nothing is invented: every code/name comes from a table cell or paragraph of the source,
and each record keeps the source file name and row/paragraph index.

Requires LibreOffice (soffice) to convert .doc -> .docx in a temporary directory, and python-docx.
Usage: python3 scripts/occupations/extract_34_2020.py
"""
import difflib
import hashlib
import json
import re
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

import docx  # python-docx

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'data/occupations/34-2020'
SRC = BASE / 'source'
LEVEL_HEADERS = ['cấp 1', 'cấp 2', 'cấp 3', 'cấp 4', 'cấp 5']


def norm(text):
    text = unicodedata.normalize('NFC', text or '')
    return re.sub(r'\s+', ' ', text).strip()


def key(text):
    return norm(text).casefold().rstrip(' .:;')


def check_manifest():
    manifest = json.loads((SRC / 'MANIFEST.json').read_text(encoding='utf-8'))
    for entry in manifest['files']:
        body = (SRC / entry['filename']).read_bytes()
        if hashlib.sha256(body).hexdigest() != entry['sha256']:
            sys.exit(f"SHA256 mismatch: {entry['filename']}")
    return sorted(entry['filename'] for entry in manifest['files'])


def convert(filenames, out_dir):
    for name in filenames:
        subprocess.run(['soffice', '--headless', '--convert-to', 'docx', '--outdir', str(out_dir),
                        str(SRC / name)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       timeout=600)
    docs = {}
    for name in filenames:
        path = out_dir / (Path(name).stem + '.docx')
        if not path.exists():
            sys.exit(f'Conversion failed: {name}')
        docs[name] = docx.Document(str(path))
    return docs


def find_code_table(document):
    for table in document.tables:
        header = [key(c.text).replace('\n', ' ') for c in table.rows[0].cells]
        header = [re.sub(r'\s+', ' ', h) for h in header]
        if header[:5] == LEVEL_HEADERS and 'tên' in header[5]:
            return table
    return None


def extract_codes(docs):
    codes, review = [], []
    found = [(name, find_code_table(d)) for name, d in docs.items()]
    found = [(n, t) for n, t in found if t is not None]
    if len(found) != 1:
        sys.exit(f'Expected exactly one Phụ lục I code table, found {len(found)}')
    source_file, table = found[0]
    last_at_level = {}
    for row_index, row in enumerate(table.rows[1:], start=1):
        cells = [norm(c.text) for c in row.cells]
        level_cells = cells[:5]
        name = cells[5]
        filled = [(i + 1, v) for i, v in enumerate(level_cells) if v]
        if not filled and not name:
            continue
        bad = [v for _, v in filled if not re.fullmatch(r'\d+', v)]
        if not filled or bad:
            review.append({'source_file': source_file, 'row': row_index, 'cells': cells,
                           'reason': 'no code cell' if not filled else 'non-numeric code cell'})
            continue
        # A row may carry a level-4 and level-5 code together with one shared name.
        wrong = [(lv, v) for lv, v in filled if len(v) != lv]
        for lv, v in wrong:   # e.g. source row 52: "1330" printed in the Cấp 5 column — not corrected
            review.append({'source_file': source_file, 'row': row_index, 'cells': cells,
                           'reason': f'code {v!r} in Cấp {lv} column has wrong length; not imported'})
        filled = [(lv, v) for lv, v in filled if len(v) == lv]
        for level, code in filled:
            parent = None
            if level > 1:
                parent = last_at_level.get(level - 1)
            codes.append({'code': code, 'level': level, 'parent_code': parent, 'name_vi': name,
                          'source_file': source_file, 'source_row': row_index})
            last_at_level[level] = code
            for deeper in range(level + 1, 6):
                last_at_level.pop(deeper, None)
    return codes, review


def paragraphs_after_appendix_ii(docs, ordered_names):
    started = False
    for name in ordered_names:
        for index, paragraph in enumerate(docs[name].paragraphs):
            text = norm(paragraph.text)
            if started is not True:
                # Explanations start at "Phần II" inside Phụ lục II (Phần I = general concepts).
                if re.fullmatch(r'phụ lục ii', key(text)):
                    started = 'appendix'
                elif started == 'appendix' and re.fullmatch(r'phần ii', key(text)):
                    started = True
                continue
            if text:
                yield name, index, text


SECTION_MARKERS = [
    ('main_tasks', re.compile(r'^nhiệm vụ chủ yếu gồm\s*:?\s*', re.I)),
    ('examples', re.compile(r'^ví dụ (về )?(các|những) nghề\b.*?:\s*', re.I)),
    ('exclusions', re.compile(r'^loại trừ\s*:?\s*', re.I)),
    ('skill_note', re.compile(r'^(?=một số nghề trong nhóm này yêu cầu)', re.I)),
]
# "2141. Name" or combined "512 - 5120. Name" (one explanation for both codes)
HEADING = re.compile(r'^(\d{1,5})(?:\s*[-–]\s*(\d{1,5}))?\.\s+(.+)$')


def extract_descriptions(docs, ordered_names, codes):
    by_code = {}
    for item in codes:
        by_code.setdefault(item['code'], item)
    descriptions, current, section = {}, None, 'description'
    name_notes, dup_headings, unlinked = [], [], []
    for source_file, index, text in paragraphs_after_appendix_ii(docs, ordered_names):
        group = re.match(r'^nhóm (\d)\.\s+(.+)$', text, re.I)   # level-1 heading: "Nhóm 2. ..."
        match = HEADING.match(f'{group.group(1)}. {group.group(2)}') if group else HEADING.match(text)
        heading_codes = [c for c in (match.group(1), match.group(2)) if c] if match else []
        if heading_codes and all(c in by_code for c in heading_codes):
            title = match.group(3)
            names = {key(by_code[c]['name_vi']) for c in heading_codes}
            exact = key(title) in names
            close = max(difflib.SequenceMatcher(None, key(title), n).ratio() for n in names)
            # A single 4-digit code heading ("7513. ...") is linked by code even when the wording
            # differs from Phụ lục I: the code fixes the occupation group (user decision 2026-10-15).
            # Fuzzy guard stays for other headings so numbered list items are never mistaken for codes.
            code_heading = len(heading_codes) == 1 and re.match(r'^\d{4}\.', text)
            if exact or close >= 0.9 or code_heading:
                official = by_code[heading_codes[-1]]['name_vi']
                record = {'code': None, 'name_vi': official,
                          'alias_vi': None if exact else norm(title),
                          'description': [], 'main_tasks': [],
                          'examples': [], 'exclusions': [], 'skill_note': [],
                          'source_file': source_file, 'source_paragraph': index,
                          'heading_codes': heading_codes}
                if not exact:
                    name_notes.append({'codes': heading_codes, 'appendix_ii_title': norm(title),
                                       'appendix_i_names': [by_code[c]['name_vi'] for c in heading_codes],
                                       'similarity': round(close, 3), 'source_file': source_file,
                                       'source_paragraph': index})
                for c in heading_codes:
                    if c in descriptions:
                        dup_headings.append({'code': c, 'source_file': source_file, 'source_paragraph': index})
                    descriptions[c] = dict(record, code=c)
                current = record
                section = 'description'
                continue
        if current is None or re.fullmatch(r'(phần [ivx]+|giải thích các nghề.*|nghề nghiệp việt nam)', key(text)):
            continue
        for name, pattern in SECTION_MARKERS:
            if pattern.match(text):
                section = name
                text = pattern.sub('', text, count=1).strip()
                break
        if text:
            current[section].append(text)
    ordered = [descriptions[c] for c in sorted(descriptions, key=lambda c: (len(c), c))]
    return ordered, name_notes, dup_headings, unlinked


def validate(codes, descriptions, review):
    from collections import Counter
    counts = Counter(c['code'] for c in codes)
    known = set(counts)
    duplicates = sorted(c for c, n in counts.items() if n > 1)
    parent_missing = [c['code'] for c in codes if c['level'] > 1 and (not c['parent_code'] or c['parent_code'] not in known)]
    parent_prefix_mismatch = [c['code'] for c in codes if c['level'] > 1 and c['parent_code'] and not c['code'].startswith(c['parent_code'])]
    level_mismatch = [c['code'] for c in codes if len(c['code']) != c['level']]
    empty_name = [c['code'] for c in codes if not c['name_vi']]
    described = {d['code'] for d in descriptions}
    per_level = {f'cap_{lv}': sum(1 for c in codes if c['level'] == lv) for lv in range(1, 6)}
    described_per_level = {f'cap_{lv}': sum(1 for c in codes if c['level'] == lv and c['code'] in described)
                           for lv in range(1, 6)}
    return {
        'counts_per_level': per_level,
        'total_codes': len(codes),
        'described_codes_per_level': described_per_level,
        'described_total': len(described),
        'not_described_total': len(known - described),
        'checks': {
            'duplicate_codes': duplicates,
            'parent_missing': parent_missing,
            'parent_prefix_mismatch': parent_prefix_mismatch,
            'level_length_mismatch': level_mismatch,
            'empty_name': empty_name,
            'needs_review_rows': len(review),
        },
    }


def main():
    filenames = check_manifest()
    with tempfile.TemporaryDirectory() as tmp:
        docs = convert(filenames, Path(tmp))
        codes, review = extract_codes(docs)
        descriptions, name_notes, dup_headings, unlinked = extract_descriptions(docs, filenames, codes)
    review += [dict(n, reason='Phụ lục II title differs from Phụ lục I name; linked by code, title kept as alias_vi') for n in name_notes]
    review += [dict(d, reason='code explained twice in Phụ lục II; last one kept') for d in dup_headings]
    report = validate(codes, descriptions, review)
    dump = lambda name, data: (BASE / name).write_text(json.dumps(data, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    dump('occupation_codes.json', codes)
    dump('descriptions.json', descriptions)
    dump('needs_review.json', review)
    dump('VALIDATION.json', report)
    print(json.dumps(report, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
